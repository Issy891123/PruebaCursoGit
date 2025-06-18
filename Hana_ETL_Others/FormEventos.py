import os
import pandas as pd
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from openpyxl import load_workbook
from hana_ml.dataframe import ConnectionContext
from datetime import datetime

# ----------------------------------------
# 1) Conexión a HANA
# ----------------------------------------
cc = ConnectionContext(
    address="548b50f4-1fe9-4725-baea-e9f96bb4f092.hana.prod-us10.hanacloud.ondemand.com",
    port=443,
    user=os.environ["DB_USER"],
    password=os.environ["DB_PASSWORD"],
    encrypt="true",
    sslValidateCertificate="false"
)

# ----------------------------------------
# 2) Validadores
# ----------------------------------------
def validate_cedula(new: str) -> bool:
    return new.isalnum() and len(new) <= 15

def validate_celular(new: str) -> bool:
    return new.isdigit() and len(new) <= 10

# ----------------------------------------
# 3) Queries exactos
# ----------------------------------------
def query_hana(tipodoc, cedula, evento, fecha):
    fullid = tipodoc + cedula

    sql1 = f"""SELECT '{evento}' AS EVENTO, '{fecha}' AS FECHA_EVENTO,
IDENTIFICADOR, TIPO_IDENTIFICACION_AFILIADO, NUM_IDENTIFICACION_AFILIADO,
PRIMER_NOMBRE, SEGUNDO_NOMBRE, PRIMER_APELLIDO, SEGUNDO_APELLIDO,
CODIGO_EXTERNO_SAP_BP, FECHA_NACIMIENTO, NIVEL_EDUCATIVO,
OCUPACION, ESTADO_CIVIL, NACIONALIDAD, TIENE_DISCAPACIDAD, EDAD,
TRABAJADOR_COLSUBSIDIO, GENERO, PROFESION, SEGMENTO_GRUPO_FAMILIAR,
GENERACIONES, NOMBRE_EMPRESA_PRINCIPAL, NUMERO_CELULAR,
CORREO_ELECTRONICO, ESTADO_AFILIACION, MARCACION_MULTIPLE_EMPRESA,
TIPO_AFILIADO, FECHA_AFILIACION, CATEGORIA_AFI, RANGO_SALARIAL,
CARGO_SOLO_RL, TIENE_BENEFICIARIO, PERIODO_GRACIA,
EN_QUE_EMPRESA_TRABAJA, SEGMENTO, FECHA_RETIRO, RANGO_EDAD,
ESTADO_EMPRESA, ID_EMPRESA, ID_PERSONA, NIVEL_SALARIO,
RAZON_SOCIAL, SEGMENTO_POBLACIONAL, TIPO_APORTANTE_EMPRESA,
NIVEL_CARGO, CLASIFICACION_CARGO, CARGO_EMPRESA,
NIVEL_DECISION, AREA_QUE_PERTENECE
FROM COLSUBSIDIO_DA.CV_PERFIL_PERSONA_FULL
WHERE IDENTIFICADOR = '{fullid}'"""
    print("[Registrar] SQL =", sql1.strip())
    df = cc.sql(sql1).collect()
    if not df.empty:
        return df, False

    sql2 = f"""SELECT '{evento}' AS EVENTO, '{fecha}' AS FECHA_EVENTO,
PC.IDENTIFICADOR_PAC,PC.TIPO_IDENTIFICACION,PC.NUMERO_IDENTIFICACION_PAC,
PC.PRIMER_NOMBRE_PERSONA_A_CARGO AS PRIMER_NOMBRE,
PC.SEGUNDO_NOMBRE_PERSONA_A_CARGO AS SEGUNDO_NOMBRE,
PC.PRIMER_APELLIDO_PAC AS PRIMER_APELLIDO,
PC.SEGUNDO_APELLIDO_PAC AS SEGUNDO_APELLIDO,
'' AS CODIGO_EXTERNO_SAP_BP,PC.FECHA_NACIMIENTO_PERSONA_A_CARGO,
PF.NIVEL_EDUCATIVO,PF.OCUPACION,PF.ESTADO_CIVIL,PF.NACIONALIDAD,
PF.TIENE_DISCAPACIDAD,PF.EDAD,PF.TRABAJADOR_COLSUBSIDIO,PF.GENERO,
PF.PROFESION,PF.SEGMENTO_GRUPO_FAMILIAR,PF.GENERACIONES,
PF.NOMBRE_EMPRESA_PRINCIPAL,PF.NUMERO_CELULAR,PF.CORREO_ELECTRONICO,
PF.ESTADO_AFILIACION,PF.MARCACION_MULTIPLE_EMPRESA,PF.TIPO_AFILIADO,
PF.FECHA_AFILIACION,PC.CATEGORIA_PERSONA_A_CARGO,PF.RANGO_SALARIAL,
PF.CARGO_SOLO_RL,PF.TIENE_BENEFICIARIO,PF.PERIODO_GRACIA,
PF.EN_QUE_EMPRESA_TRABAJA,PF.SEGMENTO,PF.FECHA_RETIRO,PF.RANGO_EDAD,
PF.ESTADO_EMPRESA,PF.ID_EMPRESA,PF.ID_PERSONA,PF.NIVEL_SALARIO,
PF.RAZON_SOCIAL,PF.SEGMENTO_POBLACIONAL,PF.TIPO_APORTANTE_EMPRESA,
PF.NIVEL_CARGO,PF.CLASIFICACION_CARGO,PF.CARGO_EMPRESA,
PF.NIVEL_DECISION,PF.AREA_QUE_PERTENECE, IDENTIFICADOR_PAC,
PC.PRIMER_NOMBRE_PERSONA_A_CARGO || ' ' || PC.SEGUNDO_NOMBRE_PERSONA_A_CARGO || ' ' ||
PC.PRIMER_APELLIDO_PAC || ' ' || PC.SEGUNDO_APELLIDO_PAC AS NOMBRE_PAC,
CASE WHEN PC.PARENTESCO IS NULL THEN 'Afiliado principal' ELSE 'Beneficiario' END AS TIPO_AFILIADO,
PC.PARENTESCO
FROM COLSUBSIDIO_DA.CV_PAC_FULL PC
INNER JOIN COLSUBSIDIO_DA.CV_IDN_PERFIL_PERSONA_FULL PF
ON PC.IDENTIFICACION_AFILIADO = PF.IDENTIFICADOR
WHERE PC.IDENTIFICADOR_PAC = '{fullid}'"""
    print("[Entró a PAC] SQL =", sql2.strip())
    df2 = cc.sql(sql2).collect()
    return df2, not df2.empty

# ----------------------------------------
# 4) Export a Excel
# ----------------------------------------
def append_to_excel(df, path, sheet="Asistentes"):
    if not os.path.exists(path):
        with pd.ExcelWriter(path, engine="openpyxl") as w:
            df.to_excel(w, sheet_name=sheet, index=False)
    else:
        wb = load_workbook(path)
        if sheet not in wb.sheetnames:
            with pd.ExcelWriter(path, engine="openpyxl", mode="a") as w:
                df.to_excel(w, sheet_name=sheet, index=False)
        else:
            ws = wb[sheet]
            start = ws.max_row
            with pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="overlay") as w:
                df.to_excel(w, sheet_name=sheet,
                            index=False, header=False,
                            startrow=start)
        wb.close()

# ----------------------------------------
# 5) Forms corregidos
# ----------------------------------------
class BaseForm(tk.Toplevel):
    def __init__(self, master, title):
        super().__init__(master)
        self.title(title)
        self.resizable(False, False)
        labels = ["Nombre", "No. de documento", "Celular", "Correo",
                  "Empresa", "Nit", "Cargo", "Interesado en"]
        self.entries = {}
        vced = self.register(validate_cedula)
        vcel = self.register(validate_celular)
        for i, txt in enumerate(labels):
            ttk.Label(self, text=txt).grid(row=i, column=0, sticky="e", padx=5, pady=2)
            cmd = (vced,"%P") if txt=="No. de documento" else (vcel,"%P") if txt=="Celular" else (self.register(lambda v: True),"%P")
            ent = ttk.Entry(self, width=40, validate="key", validatecommand=cmd)
            ent.grid(row=i, column=1, padx=5, pady=2)
            self.entries[txt] = ent
        ttk.Button(self, text="Registrar", command=self.on_register).grid(row=len(labels), column=0, columnspan=2, pady=10)
        self.result = None  # <--- aquí almacenaremos los valores

    def on_register(self):
        data = {}
        for f in ["Nombre","No. de documento","Celular","Correo","Cargo","Interesado en"]:
            val = self.entries[f].get().strip()
            if not val:
                messagebox.showwarning("Faltan datos", f"'{f}' es obligatorio.")
                self.entries[f].focus_set()
                return
            data[f] = val
        self.result = data         # <--- guardo los datos
        self.destroy()

class AfiliadoForm(BaseForm):
    def __init__(self, master, rec: dict, esPAC=False):
        title = "Beneficiario PAC" if esPAC else "Afiliado"
        super().__init__(master, title)
        # precarga
        full = rec.get("NOMBRE_PAC") or " ".join([
            rec.get("PRIMER_NOMBRE",""),
            rec.get("SEGUNDO_NOMBRE",""),
            rec.get("PRIMER_APELLIDO",""),
            rec.get("SEGUNDO_APELLIDO","")
        ]).strip()
        self.entries["Nombre"].insert(0, full)
        self.entries["Nombre"].config(state="readonly")
        doc = rec.get("IDENTIFICADOR_PAC") or rec.get("ID_PERSONA","")
        self.entries["No. de documento"].insert(0, doc)
        self.entries["No. de documento"].config(state="readonly")
        self.entries["Empresa"].insert(0, rec.get("RAZON_SOCIAL",""))
        self.entries["Empresa"].config(state="readonly")
        self.entries["Nit"].insert(0, rec.get("ID_EMPRESA",""))
        self.entries["Nit"].config(state="readonly")
        if esPAC and rec.get("PARENTESCO"):
            ttk.Label(self, text=f"Parentesco: {rec['PARENTESCO']}", foreground="blue")\
                .grid(row=0, column=2, padx=10)

class NoAfiliadoForm(BaseForm):
    def __init__(self, master, tip, ced):
        super().__init__(master, "No Afiliado")
        self.entries["No. de documento"].insert(0, tip+ced)
        self.entries["No. de documento"].config(state="readonly")

# ----------------------------------------
# 6) on_validate reasignado
# ----------------------------------------
def on_validate(master, tipvar, cedvar, ev, fe):
    tip = tipvar.get().strip()
    ced = cedvar.get().strip()
    df, esPAC = query_hana(tip, ced, ev, fe)
    if df.empty:
        form = NoAfiliadoForm(master, tip, ced)
    else:
        rec = df.iloc[0].to_dict()
        form = AfiliadoForm(master, rec, esPAC)
    form.grab_set()
    master.wait_window(form)          # <--- espero a que cierre
    if form.result is None:
        return                        # <--- el usuario canceló
    # genero el df con 1 fila y añado las columnas del form
    row = df.iloc[0].to_dict() if not df.empty else {}
    for k,v in form.result.items():
        # mapear nombres de campo
        if k=="Celular":       row["NUMERO_CELULAR"]     = v
        if k=="Correo":        row["CORREO_ELECTRONICO"] = v
        if k=="Cargo":         row["CARGO_EMPRESA"]      = v
        if k=="Interesado en": row["SERVICIO_INTERES"]   = v
    df2 = pd.DataFrame([row])
    out = os.path.join(os.environ["USERPROFILE"],"Downloads","Asistentes.xlsx")
    print(out)
    append_to_excel(df2, out)

# ----------------------------------------
# 7) Carga masiva
# ----------------------------------------
def bulk_load(master, ev, fe):
    path = filedialog.askopenfilename(filetypes=[("CSV","*.csv")], title="Seleccione CSV")
    if not path: return
    s = pd.read_csv(path, dtype=str).iloc[:,0].dropna().astype(str)
    no_aff = []
    out = os.path.join(os.environ["USERPROFILE"],"Downloads","Asistentes.xlsx")
    print(out)
    for full in s:
        tip, ced = full[:2], full[2:]
        df, esPAC = query_hana(tip, ced, ev, fe)
        if df.empty:
            no_aff.append(full)
        else:
            for col in ["NUMERO_CELULAR","CORREO_ELECTRONICO","CARGO_EMPRESA","SERVICIO_INTERES"]:
                df[col] = ""
            append_to_excel(df, out)
    if no_aff:
        messagebox.showwarning("No Afiliados", "Gestión aparte:\n"+ "\n".join(no_aff))
    else:
        messagebox.showinfo("Listo","Carga masiva completada.")

# ----------------------------------------
# 8) App principal
# ----------------------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Captura Asistentes")
        ev = tk.StringVar(value="Bon Jovi Concert")
        fe = tk.StringVar(value=datetime.now().strftime("%Y-%m-%d %H:%M"))
        ttk.Label(self, text="Evento:").grid(row=0,column=0,sticky="e")
        ttk.Entry(self, textvariable=ev).grid(row=0,column=1)
        ttk.Label(self, text="Fecha:").grid(row=1,column=0,sticky="e")
        e2=ttk.Entry(self, textvariable=fe); e2.grid(row=1,column=1); e2.config(state="readonly")
        self.tipvar=tk.StringVar(value="CC")
        self.cedvar=tk.StringVar()
        ttk.Combobox(self, textvariable=self.tipvar, values=["CC","TI","PT","CE","PA","CD","PE","RC"]).grid(row=2,column=1)
        vced=self.register(validate_cedula)
        ttk.Entry(self, textvariable=self.cedvar, validate="key", validatecommand=(vced,"%P")).grid(row=3,column=1)
        ttk.Button(self, text="Validar",
                   command=lambda: on_validate(self,self.tipvar,self.cedvar,ev.get(),fe.get())
        ).grid(row=4,column=0)
        ttk.Button(self, text="Cargar CSV", command=lambda: bulk_load(self,ev.get(),fe.get())).grid(row=4,column=1)

if __name__=="__main__":
    App().mainloop()